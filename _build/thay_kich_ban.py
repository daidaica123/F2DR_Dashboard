# -*- coding: utf-8 -*-
"""
THAY TOAN BO alert cua mot hay nhieu KICH BAN trong score_clean.

Dung khi rule duoc sua roi chay lai ca ky: du lieu cu cua kich ban do SAI
hoan toan, phai xoa het roi nap lai tu file tho moi — khac voi
bo_sung_ngay.py (chi them mot ngay con thieu).

Chay:
    py -3.10 _build/thay_kich_ban.py --topic NEW_AML_TK_moi_mo_phat_sinh...
    py -3.10 _build/thay_kich_ban.py --topic A --topic B --that

Khong truyen --topic thi liet ke cac topic co trong thu muc nguon roi dung.
"""
import argparse
import io
import os
import re
import shutil
import sys
import unicodedata
from datetime import datetime

import pandas as pd

if hasattr(sys.stdout, "buffer"):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8",
                                  errors="replace", line_buffering=True)

HERE = os.path.dirname(os.path.abspath(__file__))
GOC = os.path.dirname(HERE)
XLS = (r"D:\Project F2DR\File lưu trữ thông tin all kịch bản"
       r"\File điền thông tin kịch bản_hành vi.xlsx")
DIEM = (r"D:\Project F2DR\File lưu trữ thông tin all kịch bản"
        r"\20260625_scoring ver 3.xlsx")

ap = argparse.ArgumentParser()
ap.add_argument("--src", default=r"C:\Users\daicd\Downloads\DATA_FULL")
ap.add_argument("--csv", default=None,
                help="File score_clean dich. Khong dat = tu tim .csv.gz")
ap.add_argument("--topic", action="append", default=[],
                help="Tien to ten file cua kich ban can thay. Lap lai duoc.")
ap.add_argument("--xls", default=XLS)
ap.add_argument("--diem", default=DIEM)
ap.add_argument("--that", action="store_true")
a = ap.parse_args()

print("=" * 74)
print("THAY TOAN BO ALERT CUA KICH BAN")
print("=" * 74)


def norm(s):
    s = re.sub(r"^\[(SCORE|NEW)\]_", "", str(s or ""))
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.replace("đ", "d").replace("Đ", "D")
    return re.sub(r"[^a-z0-9]", "", s.lower())


RX_TEN = re.compile(r"^(.*?)_?(?:\d{2}_\d{2}(?:_\d)?|\d{2})?$")
fs = sorted(x for x in os.listdir(a.src) if x.lower().endswith(".csv"))
if not a.topic:
    print("Chua chon --topic. Cac topic co trong %s:" % a.src)
    for t in sorted({RX_TEN.match(f[:-4]).group(1).rstrip("_") for f in fs}):
        print("   %s" % t)
    sys.exit(0)

# ── map ten KB + bang diem ──
from openpyxl import load_workbook                                # noqa: E402

wb = load_workbook(a.xls, read_only=True, data_only=True)
rows = list(wb["FINAL"].iter_rows(values_only=True))
hdr = [str(x or "").strip() for x in rows[0]]
iKB, iTP = hdr.index("Tên Kịch bản"), hdr.index("Tên topics")
MAP = {}
for r in rows[1:]:
    kb, tp = str(r[iKB] or "").strip(), str(r[iTP] or "").strip()
    if kb.startswith("[NEW]") and tp:
        MAP[norm(tp)] = kb
wb.close()

wb = load_workbook(a.diem, read_only=True, data_only=True)
LUT = {}
for r in wb["Bảng điểm 64 KB"].iter_rows(min_row=2, values_only=True):
    ten, lv, di = r[2], r[3], r[9]
    if ten and lv and di is not None:
        LUT[norm(ten)] = (str(lv).strip().title(), int(di))
wb.close()

# Uu tien ban nen .csv.gz (nhe hon ~12 lan, pandas doc/ghi thang)
if not a.csv:
    for ten in ("score_clean_full.csv.gz", "score_clean_full.csv"):
        p = os.path.join(GOC, "data", ten)
        if os.path.exists(p):
            a.csv = p
            break
    else:
        sys.exit("Khong tim thay score_clean_full trong thu muc data")
NEN = "gzip" if a.csv.endswith(".gz") else None
d = pd.read_csv(a.csv, encoding="utf-8-sig", dtype=str)
KHUON = list(d.columns)
print("score_clean : %s dong, %d kich ban"
      % (format(len(d), ","), d.usecase_clean.nunique()))
print("-" * 74)

RX_HAU = re.compile(r"_backup\d{4}.*$")
import hashlib                                                    # noqa: E402
MUOI = os.environ.get("F2DR_SALT", "f2dr-2026-viettel-money-quan-tri-rui-ro")

giu = d
tat_ca = []
for top in a.topic:
    kb = MAP.get(norm(top))
    if not kb:
        sys.exit("Khong map duoc topic '%s' sang ten kich ban" % top)
    cu = giu[giu.usecase_clean == kb]
    cua = [f for f in fs if f.startswith(top)]
    if not cua:
        sys.exit("Khong thay file nao bat dau bang '%s'" % top)

    phan = []
    for f in cua:
        x = pd.read_csv(os.path.join(a.src, f), encoding="utf-8-sig", dtype=str)
        if not len(x):
            continue
        ma = x.get("ma_khach_hang", pd.Series([""] * len(x))).astype(str)
        ma = ma.str.replace(RX_HAU, "", regex=True)
        if "ngay_neo" in x.columns:
            g = x.ngay_neo.astype(str).str[:10]
            g = g.where(g.str.match(r"\d{4}-\d{2}-\d{2}"),
                        x.get("thoi_gian_gd", pd.Series([""] * len(x)))
                         .astype(str).str[:10])
        else:
            g = x.get("thoi_gian_gd",
                      pd.Series([""] * len(x))).astype(str).str[:10]
        gio = (x.thoi_gian_gd.astype(str).str[:19].str.replace("T", " ",
                                                               regex=False)
               if "thoi_gian_gd" in x.columns else g + " 00:00:00")
        y = pd.DataFrame({"ma": ma, "ngay": g, "gio": gio,
                          "bank": x.get("viettel_bank_code",
                                        pd.Series([""] * len(x))).astype(str)})
        phan.append(y[y.ngay.str.match(r"\d{4}-\d{2}-\d{2}", na=False)])

    t = pd.concat(phan, ignore_index=True)
    # an danh ma dang SDT (nhom KENH) — cung muoi voi luc clean lan dau
    sdt = t.ma.str.fullmatch(r"0?\d{9,10}")
    if sdt.any():
        t.loc[sdt, "ma"] = t.ma[sdt].map(
            lambda s: hashlib.sha256((MUOI + s).encode()).hexdigest()[:20])

    lv, sc = LUT[norm(kb)]
    pdate = t.ngay.str.replace("-", "", regex=False)
    ra = pd.DataFrame({
        "AlertID": pdate + "_" + t.ma,
        "CATEGORY": cu.CATEGORY.iloc[0] if len(cu) else "",
        "created_date": t.gio, "requestor": "QTRR", "updated_date": t.gio,
        "reject_reason": float("nan"),
        "usecase_name": "[SCORE]_" + kb, "status": "0", "object_key": "userId",
        "object_value": t.ma, "last_risk_score": str(float(sc)),
        "Level": lv, "Score_KB": str(sc), "viettel_bank_code": t.bank,
        "PARTITION_DATE": pdate, "ngay_alert": t.ngay, "gio_alert": t.gio,
        "usecase_clean": kb, "request_id": t.ma, "risk_level_kb": lv,
        "chi_tiet_alert": "{}",
    })[KHUON]

    print("%s" % kb)
    print("   File nguon  : %d file" % len(cua))
    print("   Cu          : %s dong, %d ngay, %s khach"
          % (format(len(cu), ","),
             cu.ngay_alert.astype(str).str[:10].nunique() if len(cu) else 0,
             format(cu.object_value.nunique(), ",") if len(cu) else 0))
    print("   Moi         : %s dong, %d ngay, %s khach"
          % (format(len(ra), ","), ra.ngay_alert.nunique(),
             format(ra.object_value.nunique(), ",")))
    print("   Chenh lech  : %+d dong" % (len(ra) - len(cu)))
    print()
    giu = giu[giu.usecase_clean != kb]          # xoa het du lieu cu cua KB
    tat_ca.append(ra)

moi = pd.concat([giu] + tat_ca, ignore_index=True)
moi = moi.sort_values(["ngay_alert", "gio_alert"], kind="stable")

# ── kiem tra ──
print("-" * 74)
loi = []
if list(moi.columns) != KHUON:
    loi.append("Thu tu cot lech")
xau = moi.object_value.astype(str)
if not xau.str.fullmatch(r"[0-9a-f]{20}").all():
    loi.append("%d ma khach khong phai 20 hex"
               % int((~xau.str.fullmatch(r"[0-9a-f]{20}")).sum()))
RX_SDT = re.compile(r"(?<![0-9a-f.])(0[35789]\d{8})(?![0-9a-f.])")
for c in moi.columns:
    if moi[c].astype(str).apply(lambda v: bool(RX_SDT.search(v))).any():
        loi.append("Cot %s con so dien thoai" % c)
for r in tat_ca:
    kb = r.usecase_clean.iloc[0]
    if int((moi.usecase_clean == kb).sum()) != len(r):
        loi.append("KB '%s' sau khi gop khong dung so dong" % kb[:40])

print("Tong  : %s -> %s dong (%+d)"
      % (format(len(d), ","), format(len(moi), ","), len(moi) - len(d)))
print("Ngay  : %d (%s -> %s)"
      % (moi.ngay_alert.astype(str).str[:10].nunique(),
         moi.ngay_alert.min(), moi.ngay_alert.max()))
print("Kich ban: %d" % moi.usecase_clean.nunique())
if loi:
    print("KHONG DAT:")
    for x in loi:
        print("   - " + x)
    sys.exit(1)
print("Kiem tra: DAT")

if not a.that:
    print()
    print("Day moi la CHAY THU. Them --that de ghi that.")
    sys.exit(0)

duoi = ".csv.gz" if NEN else ".csv"
bak = a.csv[:-len(duoi)] + "_truoc_thayKB_%s%s" % (
    datetime.now().strftime("%Y%m%d_%H%M%S"), duoi)
shutil.copy2(a.csv, bak)
print("Sao luu : %s" % os.path.basename(bak))
moi.to_csv(a.csv, index=False, encoding="utf-8-sig", compression=NEN)
print("XONG    : %s  (%s bytes)"
      % (os.path.basename(a.csv), format(os.path.getsize(a.csv), ",")))
