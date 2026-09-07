# -*- coding: utf-8 -*-
"""
CLEAN thu muc DATA_FULL (alert tho cua nhieu kich ban, nhieu ky) roi GOP
voi score_clean dang co -> score_clean_full.csv.

File tho co 33 kieu cau truc cot khac nhau; chi 4 cot la CHUNG:
    ma_khach_hang · request_id · thoi_gian_gd · viettel_bank_code
Script chi lay 4 cot do + ngay, dung lai du 21 cot theo khuon score_clean.

NGAY ALERT lay tu `ngay_neo` (ngay RULE CHAY) chu khong phai `thoi_gian_gd`
(ngay GIAO DICH XAY RA). Hai cai thuong trung nhau, nhung 2 kich ban VAY
nhin lai lich su: giao dich tu thang 6 ma rule 01/07 moi bat duoc — alert
do thuoc ve 01/07. File nao khong co ngay_neo thi lay thoi_gian_gd.

chi_tiet_alert de rong {} (chu du an chot): cac cot nghiep vu rat khac nhau
giua 33 kieu, va nhieu cot chua thong tin dinh danh.

Chay:
    py -3.10 _build/clean_data_full.py                    # chay thu
    py -3.10 _build/clean_data_full.py --that             # ghi that
"""
import argparse
import collections
import io
import os
import re
import sys
import unicodedata

import pandas as pd

if hasattr(sys.stdout, "buffer"):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8",
                                  errors="replace", line_buffering=True)

HERE = os.path.dirname(os.path.abspath(__file__))
GOC = os.path.dirname(HERE)
XLS = (r"D:\Project F2DR\File lưu trữ thông tin all kịch bản"
       r"\File điền thông tin kịch bản_hành vi.xlsx")
DIEM = (r"D:\Project F2DR\File lưu trữ thông tin all kịch bản"
        r"\20260625_scoring ver 3.xlsx")

ap = argparse.ArgumentParser()
ap.add_argument("--src", default=r"C:\Users\daicd\Downloads\DATA_FULL")
ap.add_argument("--cu", default=os.path.join(GOC, "data",
                                             "score_clean_2608_0109.csv"))
ap.add_argument("--out", default=os.path.join(GOC, "data",
                                              "score_clean_full.csv"))
ap.add_argument("--xls", default=XLS)
ap.add_argument("--diem", default=DIEM)
ap.add_argument("--tu", default="2026-07-01", help="Bo alert truoc ngay nay")
ap.add_argument("--that", action="store_true", help="Ghi that")
a = ap.parse_args()

print("=" * 74)
print("CLEAN DATA_FULL -> GOP VAO score_clean")
print("=" * 74)


def norm(s):
    s = re.sub(r"^\[(SCORE|NEW)\]_", "", str(s or ""))
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.replace("đ", "d").replace("Đ", "D")
    return re.sub(r"[^a-z0-9]", "", s.lower())


# ══════ 1. Map topic -> ten KB (sheet FINAL) ══════
from openpyxl import load_workbook                                # noqa: E402

wb = load_workbook(a.xls, read_only=True, data_only=True)
rows = list(wb["FINAL"].iter_rows(values_only=True))
hdr = [str(x or "").strip() for x in rows[0]]
iKB, iTP = hdr.index("Tên Kịch bản"), hdr.index("Tên topics")
MAP = {}
for r in rows[1:]:
    kb, tp = str(r[iKB] or "").strip(), str(r[iTP] or "").strip()
    if kb.startswith("[NEW]") and tp:
        MAP[norm(tp)] = kb
wb.close()
print("sheet FINAL : %d topic -> ten KB" % len(MAP))

# ══════ 2. Level + Score_KB (bang diem) ══════
wb = load_workbook(a.diem, read_only=True, data_only=True)
ws = wb["Bảng điểm 64 KB"]
LUT = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    ten, lv, di = r[2], r[3], r[9]
    if ten and lv and di is not None:
        LUT[norm(ten)] = (str(lv).strip().title(), int(di))
wb.close()
print("bang diem   : %d KB co Level + Score_KB" % len(LUT))

# ══════ 3. Doc file cu de lay khuon + CATEGORY ══════
cu = pd.read_csv(a.cu, encoding="utf-8-sig", dtype=str)
print("score_clean : %s dong, %d kich ban"
      % (format(len(cu), ","), cu.usecase_clean.nunique()))
KHUON = list(cu.columns)
# CATEGORY theo tung kich ban, lay tu chinh du lieu cu
CAT = cu.groupby("usecase_clean").CATEGORY.first().to_dict()
CAT_N = {norm(k): v for k, v in CAT.items()}
# nhom cua KB moi: tra tu sheet FINAL
iNH = hdr.index("Nhóm kịch bản")
wb = load_workbook(a.xls, read_only=True, data_only=True)
for r in list(wb["FINAL"].iter_rows(values_only=True))[1:]:
    kb, nh = str(r[iKB] or "").strip(), str(r[iNH] or "").strip()
    if kb.startswith("[NEW]") and nh:
        CAT_N.setdefault(norm(kb), nh)
wb.close()

# ══════ 4. Doc 74 file tho ══════
fs = sorted(x for x in os.listdir(a.src) if x.lower().endswith(".csv"))
print("DATA_FULL   : %d file" % len(fs))
print("-" * 74)

RX_TEN = re.compile(r"^(.*?)_?(?:\d{2}_\d{2}(?:_\d)?|\d{2})?$")
# KHONG neo $: co dong hau to nam giua chuoi
# (0001...18_backup0107_MA_TINH_KHONG_TON_TAI). Cat tu _backup den het.
RX_HAU = re.compile(r"_backup\d{4}.*$")
phan, thieu, bo_ngay, bo_tho = [], collections.Counter(), 0, 0

for f in fs:
    top = RX_TEN.match(f[:-4]).group(1).rstrip("_")
    kb = MAP.get(norm(top))
    if not kb:
        thieu[top] += 1
        continue
    d = pd.read_csv(os.path.join(a.src, f), encoding="utf-8-sig", dtype=str)
    if not len(d):
        continue
    bo_tho += len(d)

    # MA KHACH HANG = ma_khach_hang, KHONG phai request_id.
    # Trong bo du lieu nay request_id la MA GIAO DICH (vd 260701589359593),
    # moi giao dich mot ma khac nhau — lay no lam object_value thi dashboard
    # dem moi giao dich thanh mot khach rieng (do duoc: 233.488 "khach").
    #
    # Hau to _backupDDMM = ngay RULE CHAY, phai bo. No nam o hai vi tri:
    #   cuoi chuoi          0001...d6_backup0107
    #   GIUA chuoi          0001...18_backup0107_MA_TINH_KHONG_TON_TAI
    # nen regex khong duoc neo $.
    ma = d.get("ma_khach_hang", pd.Series([""] * len(d))).astype(str)
    ma = ma.str.replace(RX_HAU, "", regex=True)

    # NGAY: ngay_neo (ngay rule chay) truoc, khong co thi thoi_gian_gd
    if "ngay_neo" in d.columns:
        g = d.ngay_neo.astype(str).str[:10]
        g = g.where(g.str.match(r"\d{4}-\d{2}-\d{2}"),
                    d.get("thoi_gian_gd", pd.Series([""] * len(d)))
                     .astype(str).str[:10])
    else:
        g = d.get("thoi_gian_gd", pd.Series([""] * len(d))).astype(str).str[:10]

    # gio: lay tu thoi_gian_gd neu co, khong thi 00:00:00 cua ngay neo
    if "thoi_gian_gd" in d.columns:
        gio = (d.thoi_gian_gd.astype(str).str[:19].str.replace("T", " ",
                                                               regex=False))
    else:
        gio = g + " 00:00:00"

    x = pd.DataFrame({
        "kb": kb, "ma": ma, "ngay": g, "gio": gio,
        "bank": d.get("viettel_bank_code", pd.Series([""] * len(d))).astype(str),
    })
    x = x[x.ngay.str.match(r"\d{4}-\d{2}-\d{2}", na=False)]
    truoc = len(x)
    x = x[x.ngay >= a.tu]                       # bo alert truoc moc
    bo_ngay += truoc - len(x)
    if len(x):
        phan.append(x)

if thieu:
    print("!! %d topic KHONG map duoc ten KB:" % len(thieu))
    for t, n in thieu.items():
        print("   %-58s (%d file)" % (t[:58], n))
    sys.exit("Dung lai — bo sung ten topics vao sheet FINAL roi chay lai.")

t = pd.concat(phan, ignore_index=True)
print("Doc xong    : %s dong tho" % format(bo_tho, ","))
print("Bo truoc %s: %s dong (giao dich cu, rule chay sau moi bat)"
      % (a.tu, format(bo_ngay, ",")))
print("Con lai     : %s dong" % format(len(t), ","))

# ══════ 4b. An danh ma khach dang SO DIEN THOAI ══════
# Rieng nhom KENH, ma khach von la SDT that. Repo public nen phai bam.
# Bam sha256(muoi + so)[:20] -> cung dinh dang voi cac ma bam san co, nen
# dashboard khong phai sua gi. Cung mot so luon ra cung mot ma, cac ky
# khac nhau van nhan ra la mot kenh.
import hashlib                                                    # noqa: E402

MUOI = os.environ.get("F2DR_SALT", "f2dr-2026-viettel-money-quan-tri-rui-ro")
la_sdt = t.ma.astype(str).str.fullmatch(r"0?\d{9,10}")
n_sdt = int(la_sdt.sum())
if n_sdt:
    t.loc[la_sdt, "ma"] = t.ma[la_sdt].map(
        lambda s: hashlib.sha256((MUOI + s).encode()).hexdigest()[:20])
    print("An danh SDT : %s dong (%s so rieng biet) — nhom KENH"
          % (format(n_sdt, ","), format(int(la_sdt.sum() and
             t.ma[la_sdt].nunique()), ",")))

# ══════ 5. Level + Score_KB cho tung KB ══════
kbs = sorted(t.kb.unique())
print("-" * 74)
print("%d kich ban trong DATA_FULL" % len(kbs))
sot = [k for k in kbs if norm(k) not in LUT]
if sot:
    print("!! %d KB khong co trong bang diem:" % len(sot))
    for k in sot:
        print("   %s" % k)
    sys.exit("Dung lai — thieu Level/Score_KB.")
moi = [k for k in kbs if k not in set(cu.usecase_clean)]
if moi:
    print("KB MOI (chua co trong score_clean): %d" % len(moi))
    for k in moi:
        lv, sc = LUT[norm(k)]
        print("   %-56s %s / %s diem" % (k[:56], lv, sc))

# ══════ 6. Dung du 21 cot ══════
lv = t.kb.map(lambda k: LUT[norm(k)][0])
sc = t.kb.map(lambda k: LUT[norm(k)][1])
pdate = t.ngay.str.replace("-", "", regex=False)

ra = pd.DataFrame({
    "AlertID": pdate + "_" + t.ma,
    "CATEGORY": t.kb.map(lambda k: CAT.get(k) or CAT_N.get(norm(k)) or ""),
    "created_date": t.gio,
    "requestor": "QTRR",
    "updated_date": t.gio,
    "reject_reason": float("nan"),
    "usecase_name": "[SCORE]_" + t.kb,
    "status": "0",
    "object_key": "userId",
    "object_value": t.ma,
    # Hai cot nay dashboard KHONG doc (CLAUDE_CONTEXT §4.2) — no tinh lai
    # bang PP-D tu Score_KB + risk_level_kb. Ghi diem goc cho nhat quan.
    "last_risk_score": sc.astype(float).astype(str),
    "Level": lv,
    "Score_KB": sc.astype(str),
    "viettel_bank_code": t.bank,
    "PARTITION_DATE": pdate,
    "ngay_alert": t.ngay,
    "gio_alert": t.gio,
    "usecase_clean": t.kb,
    "request_id": t.ma,
    "risk_level_kb": lv,
    "chi_tiet_alert": "{}",
})[KHUON]

# ══════ 7. Kiem tra ══════
print("-" * 74)
loi = []
if list(ra.columns) != KHUON:
    loi.append("Thu tu cot khong khop")
if ra.object_value.astype(str).str.contains("_backup").any():
    loi.append("Con hau to _backup")
# Sau buoc an danh, MOI ma khach phai la hex 20 ky tu. Con dang khac tuc la
# hoac lay nham cot (request_id = ma giao dich), hoac sot SDT chua bam.
xau = ra.object_value.astype(str)
la = ~xau.str.fullmatch(r"[0-9a-f]{20}")
if la.any():
    loi.append("%s ma khach hang khong phai 20 ky tu hex (vd: %s)"
               % (format(int(la.sum()), ","), ", ".join(xau[la].head(2))))
# So khach phai hop ly so voi so alert. Lay nham cot ma giao dich thi ty le
# nay ~1.0 (moi giao dich mot "khach").
ty = ra.object_value.nunique() / max(len(ra), 1)
if ty > 0.8:
    loi.append("Ty le khach/alert = %.2f — nghi lay nham cot ma giao dich" % ty)
if (ra.CATEGORY == "").any():
    loi.append("%d dong thieu CATEGORY" % int((ra.CATEGORY == "").sum()))
RX_SDT = re.compile(r"(?<![0-9a-f.])(0[35789]\d{8})(?![0-9a-f.])")
RX_NGUOI = re.compile(r"^[A-ZÀ-Ỹ][a-zà-ỹ]+(\s+[A-ZÀ-Ỹ][a-zà-ỹ]*){1,5}$")
for c in ra.columns:
    s = ra[c].astype(str)
    if s.apply(lambda v: bool(RX_SDT.search(v))).any():
        loi.append("Cot %s con so dien thoai" % c)
    if s.apply(lambda v: bool(RX_NGUOI.match(v))).any():
        loi.append("Cot %s con ten nguoi" % c)

nd = ra.ngay_alert.nunique()
print("Dung duoc   : %s dong x %d cot" % (format(len(ra), ","), len(ra.columns)))
print("Ngay        : %d (%s -> %s)" % (nd, ra.ngay_alert.min(),
                                       ra.ngay_alert.max()))
print("Kich ban    : %d" % ra.usecase_clean.nunique())
print("Khach hang  : %s" % format(ra.object_value.nunique(), ","))
if loi:
    print("-" * 74)
    print("KHONG DAT:")
    for x in loi:
        print("   - " + x)
    sys.exit(1)
print("Kiem tra    : DAT (dung cot, khong ten, khong SDT)")

# ══════ 8. Gop voi du lieu cu ══════
print("-" * 74)
chong = set(ra.ngay_alert) & set(cu.ngay_alert.astype(str).str[:10])
if chong:
    print("!! %d ngay TRUNG voi du lieu cu: %s"
          % (len(chong), ", ".join(sorted(chong)[:5])))
    print("   -> se bi nhan doi alert. Dung lai.")
    sys.exit(1)

gop = pd.concat([cu, ra], ignore_index=True)
gop = gop.sort_values(["ngay_alert", "gio_alert"], kind="stable")
print("Gop         : %s + %s = %s dong"
      % (format(len(cu), ","), format(len(ra), ","), format(len(gop), ",")))
print("Ky moi      : %s -> %s (%d ngay)"
      % (gop.ngay_alert.min(), gop.ngay_alert.max(),
         gop.ngay_alert.astype(str).str[:10].nunique()))
print("Kich ban    : %d" % gop.usecase_clean.nunique())

if not a.that:
    print()
    print("Day moi la CHAY THU. Them --that de ghi ra %s"
          % os.path.basename(a.out))
    sys.exit(0)

gop.to_csv(a.out, index=False, encoding="utf-8-sig")
print("-" * 74)
print("XONG: %s  (%s bytes)"
      % (os.path.basename(a.out), format(os.path.getsize(a.out), ",")))
