# -*- coding: utf-8 -*-
"""
NAP ALERT HANG NGAY — append vao score_clean_full.csv.

Dung cho file alert xuat thang tu he thong (13 cot, da qua cham diem), KHAC
voi file tho cua tung kich ban (clean_data_full.py xu ly loai do).

    py -3.10 _build/nap_alert_ngay.py <file.csv>                  # chay thu
    py -3.10 _build/nap_alert_ngay.py <file.csv> --that           # ghi that
    py -3.10 _build/nap_alert_ngay.py <file.csv> --ngay 20260904 --ngay 20260905

NGAY ALERT = 8 so dau cua AlertID TRU 2 NGAY.

Vi sao tru 2: file la ngay LEN ALERT, khong phai ngay alert xay ra. Rule chay
lui 2 ngay nen alert sinh ngay 04/09 la cua giao dich ngay 02/09. Chu du an
chot 07/09/2026: "ngay 2/9 o html chinh la ngay 4/9 o file toi gui".
Doi so --lui de chinh neu do tre thay doi.

DIEM: KHONG doc last_risk_score cua file — cot do tinh theo tham so khac.
Tinh lai bang PP-D 2 tang tu Score_KB + Level cua bang diem, GOM THEO NGAY
(mot khach trong mot ngay la mot luot), y het clean_2608_0109.py va
dashboard. Xem CLAUDE_CONTEXT §4.2.
"""
import argparse
import hashlib
import io
import os
import re
import shutil
import sys
import unicodedata
from collections import defaultdict
from datetime import datetime

import pandas as pd

if hasattr(sys.stdout, "buffer"):
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8",
                                  errors="replace", line_buffering=True)

HERE = os.path.dirname(os.path.abspath(__file__))
GOC = os.path.dirname(HERE)
XLS = (r"D:\Project F2DR\File lưu trữ thông tin all kịch bản"
       r"\File điền thông tin kịch bản_hành vi.xlsx")
DIEM = (r"D:\Project F2DR\File lưu trữ thông tin all kịch bản"
        r"\20260625_scoring ver 3.xlsx")

CAP = {"Low": 25, "Medium": 50, "High": 80, "Very High": 100}
R0 = K0 = 0.7          # cung tham so voi dashboard (build_van_hanh.py)

ap = argparse.ArgumentParser()
ap.add_argument("file", help="File alert xuat tu he thong")
ap.add_argument("--csv", default=None,
                help="File score_clean dich. Khong dat thi tu tim "
                     "score_clean_full.csv.gz roi den .csv")
ap.add_argument("--ngay", action="append", default=[],
                help="Chi nap dung ngay nay — ghi theo ngay TRONG FILE "
                     "(YYYYMMDD, truoc khi tru). Lap lai duoc. Khong dat = "
                     "nap moi ngay chua co trong file dich.")
ap.add_argument("--lui", type=int, default=2,
                help="So ngay tru di de ra ngay alert that (mac dinh 2). "
                     "File ghi ngay LEN alert, rule chay lui may ngay thi "
                     "dat bang so do.")
ap.add_argument("--xls", default=XLS)
ap.add_argument("--diem", default=DIEM)
ap.add_argument("--that", action="store_true")
a = ap.parse_args()

print("=" * 74)
print("NAP ALERT HANG NGAY")
print("=" * 74)
# File dich: uu tien ban nen .csv.gz (nhe hon ~12 lan, pandas doc/ghi thang)
if not a.csv:
    for ten in ("score_clean_full.csv.gz", "score_clean_full.csv"):
        p = os.path.join(GOC, "data", ten)
        if os.path.exists(p):
            a.csv = p
            break
    else:
        sys.exit("Khong tim thay score_clean_full trong thu muc data")
for p in (a.file, a.csv):
    if not os.path.exists(p):
        sys.exit("THIEU FILE: " + p)
NEN = "gzip" if a.csv.endswith(".gz") else None


def norm(s):
    s = re.sub(r"^\[(SCORE|NEW)\]_", "", str(s or ""))
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.replace("đ", "d").replace("Đ", "D")
    return re.sub(r"[^a-z0-9]", "", s.lower())


# ══════ 1. Doc file alert ══════
d = pd.read_csv(a.file, encoding="utf-8-sig", dtype=str)
print("File nguon : %s" % os.path.basename(a.file))
print("             %s dong, %d cot" % (format(len(d), ","), len(d.columns)))
can = ["AlertID", "usecase_name", "object_value"]
thieu = [c for c in can if c not in d.columns]
if thieu:
    sys.exit("File thieu cot: " + ", ".join(thieu))

# Ngay GHI TRONG FILE = 8 so dau AlertID
d["_file"] = d.AlertID.astype(str).str[:8]
xau = ~d._file.str.fullmatch(r"20\d{6}")
if xau.any():
    sys.exit("%d AlertID khong bat dau bang YYYYMMDD (vd: %s)"
             % (int(xau.sum()), d.AlertID[xau].iloc[0]))
# Ngay ALERT THAT = lui lai a.lui ngay. Dung Timedelta chu khong tru so:
# 20260901 - 2 khong ra 20260830.
d["_ngay"] = (pd.to_datetime(d._file, format="%Y%m%d")
              - pd.Timedelta(days=a.lui)).dt.strftime("%Y%m%d")

cu = pd.read_csv(a.csv, encoding="utf-8-sig", dtype=str)
KHUON = list(cu.columns)
da_co = set(cu.ngay_alert.astype(str).str[:10])
print("File dich  : %s dong, %s -> %s"
      % (format(len(cu), ","), min(da_co), max(da_co)))
print("-" * 74)

# ══════ 2. Chon ngay can nap ══════
def gach(g):
    return "%s-%s-%s" % (g[:4], g[4:6], g[6:])


co_file = sorted(d._file.unique())
print("Ngay trong file nguon (lui %d ngay ra ngay alert that):" % a.lui)
print("   %-12s %-12s %8s   %s" % ("TRONG FILE", "-> ALERT", "DONG", ""))
for g in co_file:
    ng = d._ngay[d._file == g].iloc[0]
    n = int((d._file == g).sum())
    print("   %-12s %-12s %8s   %s"
          % (gach(g), gach(ng), format(n, ","),
             "DA CO trong file dich" if gach(ng) in da_co else ""))
if a.ngay:
    chon = [g for g in co_file if g in set(a.ngay)]
    la = set(a.ngay) - set(co_file)
    if la:
        sys.exit("Ngay %s khong co trong file nguon" % ", ".join(sorted(la)))
else:
    chon = [g for g in co_file
            if gach(d._ngay[d._file == g].iloc[0]) not in da_co]
if not chon:
    print()
    print("Khong co ngay nao moi de nap. Dung.")
    sys.exit(0)
d = d[d._file.isin(chon)].copy()
ra_ngay = sorted(d._ngay.unique())
print()
print("SE NAP : %s  ->  ngay alert %s"
      % (", ".join(gach(g) for g in chon), ", ".join(gach(g) for g in ra_ngay)))
print("         %s dong" % format(len(d), ","))

# chan nap de len ngay da co
trung = {gach(g) for g in ra_ngay} & da_co
if trung:
    sys.exit("Ngay alert %s DA CO trong file dich -> se nhan doi. Dung lai. "
             "Muon thay thi xoa ngay do truoc." % ", ".join(sorted(trung)))

# ══════ 3. Ten KB + Level + Score_KB ══════
from openpyxl import load_workbook                                # noqa: E402

wb = load_workbook(a.diem, read_only=True, data_only=True)
LUT = {}
for r in wb["Bảng điểm 64 KB"].iter_rows(min_row=2, values_only=True):
    ten, lv, di = r[2], r[3], r[9]
    if ten and lv and di is not None:
        LUT[norm(ten)] = (str(lv).strip().title(), int(di))
wb.close()
print("-" * 74)
print("Bang diem  : %d KB" % len(LUT))

# usecase_name da la ten KB day du, chi bo tien to [SCORE]_
d["_kb"] = d.usecase_name.astype(str).str.replace(r"^\[SCORE\]_", "",
                                                  regex=True)
sot = sorted({k for k in d._kb.unique() if norm(k) not in LUT})
if sot:
    print("!! %d kich ban KHONG co trong bang diem:" % len(sot))
    for k in sot:
        print("   %s" % k)
    sys.exit("Dung lai — bo sung vao bang diem roi chay lai.")

moi = sorted(set(d._kb) - set(cu.usecase_clean))
if moi:
    print("KB MOI (chua tung co trong file dich): %d" % len(moi))
    for k in moi:
        lv, sc = LUT[norm(k)]
        print("   %-56s %s / %s diem" % (k[:56], lv, sc))

# ══════ 4. CATEGORY: chuan hoa theo ten nhom dang dung ══════
# File alert ghi 'NEW_AML', 'NEW_FRESO', 'CHUYEN/ NHAN TIEN' (thua dau cach)
# trong khi file dich dung 'AML', 'SAN FRESO', 'CHUYEN/NHAN TIEN'. De nguyen
# thi mot nhom bi tach doi tren dashboard.
CAT_CU = cu.groupby("usecase_clean").CATEGORY.first().to_dict()
CAT_N = {norm(k): v for k, v in CAT_CU.items()}
wb = load_workbook(a.xls, read_only=True, data_only=True)
rows = list(wb["FINAL"].iter_rows(values_only=True))
hdr = [str(x or "").strip() for x in rows[0]]
iKB, iNH = hdr.index("Tên Kịch bản"), hdr.index("Nhóm kịch bản")
for r in rows[1:]:
    kb, nh = str(r[iKB] or "").strip(), str(r[iNH] or "").strip()
    if kb.startswith("[NEW]") and nh:
        CAT_N.setdefault(norm(kb), nh)
wb.close()
d["_cat"] = d._kb.map(lambda k: CAT_N.get(norm(k), ""))
if (d._cat == "").any():
    la = sorted(d._kb[d._cat == ""].unique())
    sys.exit("Khong tra duoc nhom nghiep vu cho: %s" % ", ".join(la[:3]))
doi = {}
for goc, chuan in zip(d.CATEGORY.astype(str), d._cat):
    if goc != chuan:
        doi[goc] = chuan
if doi:
    print("Chuan hoa ten nhom:")
    for k, v in doi.items():
        print("   %-24s -> %s" % (k, v))

# ══════ 5. An danh ma khach dang SDT ══════
MUOI = os.environ.get("F2DR_SALT", "f2dr-2026-viettel-money-quan-tri-rui-ro")
ma = d.object_value.astype(str)
sdt = ma.str.fullmatch(r"0?\d{9,10}")
if sdt.any():
    ma = ma.copy()
    ma[sdt] = ma[sdt].map(
        lambda s: hashlib.sha256((MUOI + s).encode()).hexdigest()[:20])
    print("An danh SDT: %s dong (%s so rieng biet)"
          % (format(int(sdt.sum()), ","),
             format(d.object_value[sdt].nunique(), ",")))
d["_ma"] = ma
la = ~d._ma.str.fullmatch(r"[0-9a-f]{20}")
if la.any():
    sys.exit("%d ma khach khong phai 20 hex sau khi an danh (vd: %s)"
             % (int(la.sum()), d._ma[la].iloc[0]))

# ══════ 6. Tinh SCORE NEW — gom theo NGAY ══════
# Mot LUOT = mot khach trong mot ngay. Phai gom truoc roi moi tinh diem,
# neu tinh tung dong thi moi alert thanh mot diem rieng.
def eff(base, cap, n, r):
    return base if cap <= base else cap - (cap - base) * (r ** (n - 1))


luot = defaultdict(lambda: defaultdict(int))
# Duyet bang zip chu khong itertuples: itertuples doi ten cot bat dau bang
# dau gach duoi thanh _1, _2... nen r._ma khong ton tai.
for m, g, k in zip(d._ma, d._ngay, d._kb):
    luot[(m, g)][k] += 1
diem = {}
for key, kbs in luot.items():
    es = []
    for k, n in kbs.items():
        lv, sc = LUT[norm(k)]
        es.append(eff(sc, CAP[lv], n, R0))
    M = max(es)
    prod, used = 1.0, False
    for e in es:
        if not used and e == M:
            used = True
            continue
        prod *= (1 - e / 100)
    diem[key] = min(100.0, max(0.0, M + (100 - M) * K0 * (1 - prod)))
print("-" * 74)
print("Luot (KH x ngay): %s" % format(len(luot), ","))

# ══════ 7. Dung du 21 cot ══════
gio = d.created_date.astype(str).str[:19] if "created_date" in d.columns \
    else d._ngay + " 00:00:00"
lv = d._kb.map(lambda k: LUT[norm(k)][0])
sc = d._kb.map(lambda k: LUT[norm(k)][1])

ra = pd.DataFrame({
    "AlertID": d.AlertID.astype(str),
    "CATEGORY": d._cat,
    "created_date": gio,
    "requestor": d.get("requestor", "QTRR"),
    "updated_date": d.get("updated_date", gio),
    # File nguon nhet MO TA KICH BAN vao reject_reason; file dich de trong.
    "reject_reason": float("nan"),
    "usecase_name": "[SCORE]_" + d._kb,
    "status": d.get("status", "0"),
    "object_key": d.get("object_key", "userId"),
    "object_value": d._ma,
    "last_risk_score": [("%.10g" % diem[(m, g)])
                        for m, g in zip(d._ma, d._ngay)],
    "Level": lv, "Score_KB": sc.astype(str),
    "viettel_bank_code": d.get("viettel_bank_code", "").fillna(""),
    "PARTITION_DATE": d._ngay,
    "ngay_alert": d._ngay.str[:4] + "-" + d._ngay.str[4:6] + "-" + d._ngay.str[6:],
    "gio_alert": gio,
    "usecase_clean": d._kb,
    "request_id": d._ma,
    "risk_level_kb": lv,
    "chi_tiet_alert": "{}",
})[KHUON]

# ══════ 8. Kiem tra ══════
print("-" * 74)
loi = []
if list(ra.columns) != KHUON:
    loi.append("Thu tu cot lech")
if ra.object_value.astype(str).str.contains("_backup").any():
    loi.append("Con hau to _backup")
RX_SDT = re.compile(r"(?<![0-9a-f.])(0[35789]\d{8})(?![0-9a-f.])")
RX_NG = re.compile(r"^[A-ZÀ-Ỹ][a-zà-ỹ]+(\s+[A-ZÀ-Ỹ][a-zà-ỹ]*){1,5}$")
for c in ra.columns:
    s = ra[c].astype(str)
    # AlertID = YYYYMMDDHHMMSS_<so ngau nhien> do he thong sinh; phan so sau
    # dau gach TINH CO trung dang so dien thoai (1.406/12.142 dong dinh).
    # Do khong phai du lieu ca nhan nen bo qua cot nay, cac cot khac van quet.
    if c != "AlertID" and s.apply(lambda v: bool(RX_SDT.search(v))).any():
        loi.append("Cot %s con so dien thoai" % c)
    # Level chua 'Very High' — dung dang 'Chu hoa + chu thuong' nen regex ten
    # nguoi bat nham. Day la ten muc Impact, khong phai ten khach.
    if c != "Level" and s.apply(lambda v: bool(RX_NG.match(v))).any():
        loi.append("Cot %s con ten nguoi" % c)
if int(ra.ngay_alert.nunique()) != len(chon):
    loi.append("So ngay khong khop")
# Chan chac lan nua: khong duoc dung vao ngay da co trong file dich
de = set(ra.ngay_alert.unique()) & da_co
if de:
    loi.append("Ghi de len ngay da co: %s" % ", ".join(sorted(de)))

print("Dung duoc : %s dong x %d cot" % (format(len(ra), ","), len(ra.columns)))
print("Ngay      : %s" % ", ".join(sorted(ra.ngay_alert.unique())))
print("Kich ban  : %d" % ra.usecase_clean.nunique())
print("Khach     : %s" % format(ra.object_value.nunique(), ","))
print("Diem      : %.1f - %.1f" % (min(diem.values()), max(diem.values())))
if loi:
    print("KHONG DAT:")
    for x in loi:
        print("   - " + x)
    sys.exit(1)
print("Kiem tra  : DAT")

gop = pd.concat([cu, ra], ignore_index=True)
gop = gop.sort_values(["ngay_alert", "gio_alert"], kind="stable")
print("-" * 74)
print("Gop       : %s + %s = %s dong"
      % (format(len(cu), ","), format(len(ra), ","), format(len(gop), ",")))
print("Ky moi    : %s -> %s (%d ngay)"
      % (gop.ngay_alert.min(), gop.ngay_alert.max(),
         gop.ngay_alert.astype(str).str[:10].nunique()))

if not a.that:
    print()
    print("Day moi la CHAY THU. Them --that de ghi that.")
    sys.exit(0)

duoi = ".csv.gz" if NEN else ".csv"
bak = a.csv[:-len(duoi)] + "_truoc_nap_%s%s" % (
    datetime.now().strftime("%Y%m%d_%H%M%S"), duoi)
shutil.copy2(a.csv, bak)
print("Sao luu   : %s" % os.path.basename(bak))
gop.to_csv(a.csv, index=False, encoding="utf-8-sig", compression=NEN)
print("XONG      : %s  (%s bytes)"
      % (os.path.basename(a.csv), format(os.path.getsize(a.csv), ",")))
print()
print("Buoc tiep : py -3.10 _build/build_van_hanh.py")
print("            py -3.10 _build/xem_truoc.py")
